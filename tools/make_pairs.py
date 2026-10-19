"""Build the five ready-to-test sample pairs in samples/pairs/ (drawing + BoC + answer key).

Every BoC here is SYNTHETIC: rows were transcribed by hand from public NIST MBE PMI test drawings
(requirement text read off the rendered sheets, locations from the region reader's boxes for that
text), then misses were planted on purpose. None of them is a real customer BoC.

The script also runs each pair through the real pipeline, rules only (no model, no key), and writes
what the app actually flagged into the pair's ANSWER-KEY.md, so the key always matches the app.

Usage: TALLYMARK_DATA_DIR=<temp dir> python -m tools.make_pairs [--no-check]
"""
from __future__ import annotations

import argparse
import csv
import json
import shutil
import time
import zipfile
from pathlib import Path

from app import store
from app.pipeline.geometry import Box, union_box, zone_label
from app.pipeline.run import _cache_dir
from tools.make_samples import COLUMNS

ROOT = Path(__file__).resolve().parent.parent
SAMPLES = ROOT / "samples"
OUT = SAMPLES / "pairs"
RAW = "https://github.com/PedroHang/radhires-tallymark/raw/main/samples/pairs"
SYNTHETIC = ("Synthetic BoC built for testing from a public NIST test drawing; misses were planted on purpose. "
             "It is not a real customer BoC.")

# (sheet, char_type, requirement, places, region anchors). An anchor is the region id suffix the reader
# gives that callout ("o22" = region s<sheet>-o22); the row's location is the union of its anchors.
DRAWINGS = {
    "ctc03": {
        "pdf": SAMPLES / "eval" / "nist_ctc_03_asme1_rc.pdf",
        "part_number": "NIST CTC-03", "part_name": "PMI Complex Test Case 3",
        "rows": [
            (1, "dimension", "2X Ø.438 ±.005", 2, ["o3"]),
            (1, "gdt", "⌖ Ø.05 Ⓜ A B Ⓜ C Ⓜ", 2, ["o2"]),
            (1, "dimension", "Ø2.00 ±.01", 1, ["o5"]),
            (1, "gdt", "⌖ Ø.06 D B C", 1, ["o4"]),
            (1, "gdt", "⌓ .03 A B C", 1, ["o7"]),
            (1, "dimension", "4X Ø.625 ±.005", 4, ["o8"]),
            (1, "gdt", "⌖ Ø.05 Ⓜ D B C SWITCH MOUNTING LOCATIONS", 4, ["o9", "o10", "o12"]),
            (1, "gdt", "▱ .01, 2 COPLANAR SURFACES", 1, ["r33", "r34"]),
            (1, "dimension", "Ø.438 ±.005", 1, ["o11"]),
            (1, "gdt", "⌖ Ø.02 A B", 1, ["o13"]),
            (1, "dimension", "Ø1.065 ±.003", 1, ["o14"]),
            (1, "gdt", "⊥ Ø.01 E", 1, ["o15"]),
            (1, "dimension", "Ø.438 ±.005", 1, ["r36"]),
            (1, "gdt", "⊥ Ø.01 A", 1, ["o19"]),
            (1, "gdt", "∠ .04 A", 1, ["r38"]),
            (1, "gdt", "▱ .005 / .25 X .25", 1, ["r39"]),
            (1, "gdt", "⌓ .06 A B C", 1, ["o26"]),
            (1, "dimension", "4X .82 ±.06", 4, ["o22"]),
            (1, "dimension", "Ø1.500 ±.005", 1, ["o23"]),
            (1, "gdt", "⌖ Ø.08 D B C", 1, ["o25"]),
            (1, "gdt", "⌖ .03 D B C", 1, ["r37"]),
        ],
    },
    "ctc04": {
        "pdf": SAMPLES / "eval" / "nist_ctc_04_asme1_rd.pdf",
        "part_number": "NIST CTC-04", "part_name": "PMI Complex Test Case 4",
        "rows": [
            (1, "gdt", "⌓ 2 D G H 2 SURFACES", 1, ["r29", "o5"]),
            (1, "gdt", "⌓ 0.2 D 2 SURFACES", 1, ["o3", "o5"]),
            (1, "dimension", "6X Ø6.65 ±0.12", 6, ["r26"]),
            (1, "gdt", "⌖ Ø1.5 D E F", 6, ["o12"]),
            (1, "gdt", "⌖ Ø0.3 D E", 6, ["r27", "r28"]),
            (1, "thread", "4X M12 x 1.75 - 6H ↧ 25 ±0.25", 4, ["o13"]),
            (1, "gdt", "⌖ Ø0.75 Ⓟ 50 A B C", 4, ["o14"]),
            (1, "dimension", "4X ⌴ Ø20 ±0.2", 4, ["o15"]),
            (1, "dimension", "↧ 20 ±0.2", 4, ["o15"]),
            (1, "dimension", "30X Ø14 ±0.1", 30, ["o17"]),
            (1, "gdt", "⌖ Ø0.35 A B C", 1, ["o18"]),
            (1, "gdt", "⌓ 0.5 A B C", 1, ["o19"]),
            (1, "dimension", "30X ⌵ Ø20 ±0.3 X 90° ±1°", 30, ["o20"]),
        ],
    },
    "ftc11": {
        "pdf": SAMPLES / "eval" / "nist_ftc_11_asme1_rb.pdf",
        "part_number": "NIST FTC-11", "part_name": "PMI Fully-Toleranced Test Case 11",
        "rows": [
            (1, "note", "NOTE 4: NON-RIGID PART. DIRECTLY-TOLERANCED DIMENSIONS AND TOLERANCES APPLY IN FREE STATE.",
             1, ["o24", "o25", "o26"]),
            (1, "dimension", "Ø63 ±0.5 AVG (O.D.)", 1, ["o15"]),
            (1, "dimension", "Ø32 ±0.3 AVG (I.D.)", 1, ["o18"]),
            (1, "gdt", "○ 1.2 Ⓕ", 1, ["o19", "r34"]),
            (1, "gdt", "↗ 1.2 Ⓕ A B", 1, ["r36"]),
            (1, "dimension", "2X R1.5 ±0.04", 2, ["o28"]),
            (2, "gdt", "⌓ 0.3 Ⓕ A B", 1, ["o1", "o2"]),
            (2, "dimension", "2X R0.1 MAX", 2, ["o3"]),
            (2, "dimension", "1.93 ±0.08", 1, ["o4"]),
            (2, "gdt", "▱ 0.3 Ⓕ", 1, ["o5"]),
        ],
    },
}


def _regions(pdf: Path, sheet: int) -> dict[str, dict]:
    cache = _cache_dir(store.sha256(pdf.read_bytes()))
    path = cache / f"sheet_{sheet}_regions.json"
    if not store.materialize(path):
        raise SystemExit(f"No OCR cache for {pdf.name} sheet {sheet}: run a check on it once first.")
    return {r["region_id"].split("-", 1)[1]: r for r in json.loads(path.read_text(encoding="utf-8"))}


def drawing_rows(key: str) -> list[dict]:
    spec = DRAWINGS[key]
    cache: dict[int, dict] = {}
    rows = []
    for i, (sheet, ctype, req, places, anchors) in enumerate(spec["rows"], start=1):
        regs = cache.setdefault(sheet, _regions(spec["pdf"], sheet))
        missing = [a for a in anchors if a not in regs]
        if missing:
            raise SystemExit(f"{key} sheet {sheet}: anchors {missing} not in the OCR cache")
        b = union_box([Box.from_dict(regs[a]["bbox"]) for a in anchors])
        rows.append({"char_no": str(i), "sheet": sheet, "zone": zone_label(b.center), "char_type": ctype,
                     "requirement": req, "places": places,
                     "bbox_x": round(b.x, 4), "bbox_y": round(b.y, 4), "bbox_w": round(b.w, 4), "bbox_h": round(b.h, 4)})
    return rows


def ftc06_rows() -> list[dict]:
    with (SAMPLES / "ftc06_boc_groundtruth.csv").open(encoding="utf-8") as f:
        return list(csv.DictReader(f))


# Each planted change names the row by (sheet, requirement).
PAIRS = [
    {"folder": "01-ftc06-clean-form3", "drawing": "ftc06", "format": "form3",
     "title": "FTC-06, complete Form 3 BoC (nothing missing)",
     "summary": "Complete BoC: all 48 rows, no coordinates. Nothing planted.",
     "planted": []},
    {"folder": "02-ftc06-four-misses-csv", "drawing": "ftc06", "format": "csv_simulated",
     "title": "FTC-06, CSV with coordinates and 4 planted misses",
     "summary": "The bundled demo BoC: 3 rows deleted (a general-note frame, a dimension, a position frame) "
                "and one 4X row set to 1 place.",
     "planted": [
         {"op": "drop", "sheet": 1, "requirement": "⌓ .05 A B C APPLIES TO ALL UNTOLERANCED SURFACES (NOTE 4)",
          "kind": "general-note profile frame"},
         {"op": "places", "sheet": 1, "requirement": "4X Ø.281 ±.008", "to": 1, "kind": "places count"},
         {"op": "drop", "sheet": 2, "requirement": "12.00 ±.01", "kind": "toleranced dimension"},
         {"op": "drop", "sheet": 3, "requirement": "⌖ .025 D C K", "kind": "position frame"},
     ]},
    {"folder": "03-ctc04-form3-misses", "drawing": "ctc04", "format": "form3",
     "title": "CTC-04, Form 3 BoC (no coordinates) with 3 planted misses",
     "summary": "A profile frame and a 30X hole row deleted; the 4X M12 thread row set to 1 place.",
     "planted": [
         {"op": "drop", "sheet": 1, "requirement": "⌓ 0.5 A B C", "kind": "profile frame"},
         {"op": "drop", "sheet": 1, "requirement": "30X Ø14 ±0.1", "kind": "toleranced dimension"},
         {"op": "places", "sheet": 1, "requirement": "4X M12 x 1.75 - 6H ↧ 25 ±0.25", "to": 1, "kind": "places count"},
     ]},
    {"folder": "04-ctc03-csv-misses", "drawing": "ctc03", "format": "csv",
     "title": "CTC-03, CSV with coordinates and 3 planted misses",
     "summary": "A profile frame and a diameter row deleted; the 4X .82 row set to 1 place.",
     "planted": [
         {"op": "drop", "sheet": 1, "requirement": "⌓ .06 A B C", "kind": "profile frame"},
         {"op": "drop", "sheet": 1, "requirement": "Ø2.00 ±.01", "kind": "toleranced dimension"},
         {"op": "places", "sheet": 1, "requirement": "4X .82 ±.06", "to": 1, "kind": "places count"},
     ]},
    {"folder": "05-ftc11-multisheet", "drawing": "ftc11", "format": "csv",
     "title": "FTC-11, 2 sheets, CSV with coordinates, misses on both sheets",
     "summary": "A diameter row deleted on sheet 1; a profile frame deleted and the 2X R0.1 row set to 1 place on sheet 2.",
     "planted": [
         {"op": "drop", "sheet": 1, "requirement": "Ø32 ±0.3 AVG (I.D.)", "kind": "toleranced dimension"},
         {"op": "drop", "sheet": 2, "requirement": "⌓ 0.3 Ⓕ A B", "kind": "profile frame"},
         {"op": "places", "sheet": 2, "requirement": "2X R0.1 MAX", "to": 1, "kind": "places count"},
     ]},
]

FTC06 = {"pdf": SAMPLES / "nist_ftc_06_asme1_rd.pdf", "part_number": "NIST FTC-06",
         "part_name": "PMI Fully-Toleranced Test Case 6"}
FORMAT_LABEL = {"form3": "AS9102 Form 3 layout (.xlsx)", "csv": "CSV with coordinates",
                "csv_simulated": "CSV with coordinates"}


def _spec(key: str) -> dict:
    return FTC06 if key == "ftc06" else DRAWINGS[key]


def full_rows(key: str) -> list[dict]:
    if key == "ftc06":
        rows = ftc06_rows()
        for r in rows:
            r["sheet"], r["places"] = int(r["sheet"]), int(r["places"])
        return rows
    return drawing_rows(key)


def plant(rows: list[dict], planted: list[dict]) -> tuple[list[dict], list[dict]]:
    """Apply the planted changes, renumber (gaps would give the deleted rows away) and return the
    BoC rows plus one record per planted item."""
    out = [dict(r) for r in rows]
    records = []
    for p in planted:
        match = [r for r in out if r["requirement"] == p["requirement"] and int(r["sheet"]) == p["sheet"]]
        if len(match) != 1:
            raise SystemExit(f"planted {p} matched {len(match)} rows")
        r = match[0]
        rec = {**p, "orig_char_no": r["char_no"], "zone": r["zone"], "places_true": int(r["places"]),
               "bbox": {k: float(r[f"bbox_{k}"]) for k in "xywh"}}
        if p["op"] == "places":
            r["places"] = p["to"]
            rec["row"] = r
        else:
            out.remove(r)
        records.append(rec)
    for i, r in enumerate(out, start=1):
        r["char_no"] = str(i)
    for rec in records:
        if "row" in rec:
            rec["new_char_no"] = rec.pop("row")["char_no"]
    return out, records


def write_csv(path: Path, rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS, extrasaction="ignore")
        w.writeheader()
        for r in rows:
            w.writerow(r)


def write_form3(path: Path, rows: list[dict], spec: dict) -> None:
    """AS9102 Form 3 layout: title, a note, the numbered header fields, the two section bands, then
    one row per characteristic with Results left blank (this is the plan before anyone measures)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    wb = Workbook()
    ws = wb.active
    ws.title = "Form 3 - BoC"
    widths = {"A": 10, "B": 20, "C": 16, "D": 52, "E": 10, "F": 14, "G": 18, "H": 16, "I": 18, "J": 11}
    for col, wdt in widths.items():
        ws.column_dimensions[col].width = wdt
    thin = Side(style="thin", color="999999")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    bold = Font(bold=True)
    ws["A1"] = "FORM 3 - CHARACTERISTIC ACCOUNTABILITY, VERIFICATION, AND COMPATIBILITY EVALUATION (AS9102 layout)"
    ws["A1"].font = Font(bold=True, size=12)
    ws.merge_cells("A1:J1")
    ws["A2"] = (f"SAMPLE FOR TESTING TALLYMARK. The drawing is real ({spec['part_number']}, public). "
                + SYNTHETIC + " Results are blank on purpose: this is the plan before anyone measures the part.")
    ws["A2"].font = Font(italic=True, color="AA0000")
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells("A2:J2")
    ws.row_dimensions[2].height = 30
    for cell, val in (("A3", "1. Part Number:"), ("B3", spec["part_number"]), ("C3", "2. Part Name:"),
                      ("D3", spec["part_name"]), ("E3", "3. Serial Number:"), ("G3", "4. FAIR Number:")):
        ws[cell] = val
        if val.endswith(":"):
            ws[cell].font = bold
    band = PatternFill("solid", fgColor="F2F2F2")
    ws["A4"] = "Characteristic Accountability (filled when the drawing is ballooned)"
    ws["F4"] = "Inspection / Test Results (filled after measuring the first part)"
    for c in ("A4", "F4"):
        ws[c].font, ws[c].fill = bold, band
    ws.merge_cells("A4:E4")
    ws.merge_cells("F4:J4")
    headers = ["5. Char. No.", "6. Reference Location", "7. Characteristic Designator", "8. Requirement",
               "Places (qty)", "9. Results", "10. Designed / Qualified Tooling", "11. Nonconformance Number",
               "14. Comments", "Type"]
    head_fill = PatternFill("solid", fgColor="DDE4EE")
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=5, column=i, value=h)
        c.font, c.fill, c.border = bold, head_fill, box
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for n, r in enumerate(rows, start=6):
        vals = [int(r["char_no"]), f"Sheet {r['sheet']}, zone {r['zone']}", None, r["requirement"],
                int(r["places"]), None, None, None, None, r["char_type"]]
        for i, v in enumerate(vals, start=1):
            c = ws.cell(row=n, column=i, value=v)
            c.border = box
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A6"
    wb.save(path)


# ----------------------------------------------------------------------------- verification
def notices(doc: dict) -> list[str]:
    """The run notices, worded as the results screen words them (frontend/app.js runNotices)."""
    out = []
    for w in doc.get("warnings") or []:
        if w["type"] == "no_bbox":
            out.append("This BoC has no location columns. Rows were located by matching text, so a few matches may be off.")
        elif w["type"] == "rows_not_located":
            n = len(w["char_nos"])
            more = ", ..." if n > 8 else ""
            plural = "" if n == 1 else "s"
            out.append(f"{n} row{plural} could not be placed on the drawing (CN {', '.join(w['char_nos'][:8])}{more}). "
                       "Anything they cover may show as unaccounted.")
        elif w["type"] == "sheet_without_rows":
            out.append(f"Sheet {w['sheet']} has no BoC rows, so every requirement on it shows as unaccounted.")
    st = doc.get("stats") or {}
    if st.get("ambiguous") and st.get("llm_error"):
        n = st["ambiguous"]
        verb = " is" if n == 1 else "s are"
        out.append(f"The model check did not run (it is turned off). {n} ambiguous leftover{verb} flagged with "
                   "low confidence for you to review.")
    return out


def _hit(rec: dict, flags: list[dict]) -> dict | None:
    b = Box.from_dict(rec["bbox"])
    want = "places" if rec["op"] == "places" else "unaccounted"
    for f in flags:
        if f["sheet"] == rec["sheet"] and f["class"] == want and Box.from_dict(f["bbox"]).pad(0.01).contains(b.center):
            return f
    return None


def check_pair(drawing: Path, boc: Path) -> tuple[dict, float]:
    from tools.check import run_check

    t0 = time.perf_counter()
    doc = run_check(drawing, boc, use_llm=False)
    return doc, time.perf_counter() - t0


def _md_cell(text: str) -> str:
    return str(text).replace("|", "\\|")


def answer_key(pair: dict, spec: dict, rows: list[dict], full: list[dict], records: list[dict], boc_name: str,
               doc: dict | None, secs: float | None, today: str, full_flags: int | None = None) -> tuple[str, bool]:
    pdf = spec["pdf"]
    sheets = doc["drawing"]["sheets"] if doc else max(int(r["sheet"]) for r in full)
    fmt = FORMAT_LABEL[pair["format"]]
    loc = ("no coordinates: each row's Reference Location gives the sheet and zone, and Tallymark places the row "
           "by matching its text" if pair["format"] == "form3" else "the bbox columns give each row's box on its sheet (page fractions)")
    lines = [f"# Answer key: {pair['title']}", "",
             f"> **{SYNTHETIC}**", "",
             "| | |", "| --- | --- |",
             f"| Drawing | `drawing.pdf`: {spec['part_number']} ({spec['part_name']}), {sheets} sheet{'s' if sheets > 1 else ''}, "
             f"from the public NIST MBE PMI test set (`{pdf.name}`) |",
             f"| BoC | `{boc_name}`: {fmt}, {len(rows)} rows, {loc} |",
             f"| The complete BoC | {len(full)} rows, transcribed by hand from the rendered sheets |",
             f"| Planted | {len(records) if records else 'nothing (this BoC is complete)'} |", "",
             "## How to test", "",
             "1. Open Tallymark ([live](https://radhires-tallymark.vercel.app/) or `http://localhost:8031`).",
             f"2. Drop `drawing.pdf` on the drawing box and `{boc_name}` on the BoC box.",
             "3. Leave **Use the AI** off (the answer below is the rules-only check), then click **Run check**.",
             "4. Compare the flag list on the right with the table below.", ""]
    hits: dict[int, dict] = {}
    matched = True
    if doc:
        free = list(doc["flags"])
        for i, rec in enumerate(records):
            f = _hit(rec, free)
            if f:
                hits[i] = f
                free.remove(f)
        matched = len(hits) == len(records) and not free and doc["status"] == "complete"
    if not records:
        lines += ["## Expected flags: none", "",
                  "Every requirement on the drawing has a row, so the results screen should say nothing is "
                  "missing and show **0 flags**. Covered items are greyed out on the drawing.", ""]
    else:
        lines += [f"## Expected flags: {len(records)}", "",
                  "| Flag | Type | Sheet | Zone | Requirement (as in the complete BoC) | Text the app read on the drawing | Why it is flagged |",
                  "| --- | --- | --- | --- | --- | --- | --- |"]
        order = {f["flag_id"]: n for n, f in enumerate(doc["flags"])} if doc else {}
        for i, rec in sorted(enumerate(records), key=lambda ir: order.get((hits.get(ir[0]) or {}).get("flag_id"), 99)):
            f = hits.get(i)
            fid = f["flag_id"] if f else "(not found)"
            typ = "P (places short)" if rec["op"] == "places" else "U (unaccounted)"
            zone = f["zone"] if f else rec["zone"]
            if rec["op"] == "places":
                why = (f"{rec['kind']}: row CN {rec['new_char_no']} says Places {rec['to']}, the callout says "
                       f"{rec['places_true']}X (complete BoC: {rec['places_true']})")
            else:
                why = f"{rec['kind']}: row {rec['orig_char_no']} of the complete BoC was deleted"
            read = f"`{_md_cell(f['drawing_text'])}`" if f else ""
            lines.append(f"| {fid} | {typ} | {rec['sheet']} | {zone} | {_md_cell(rec['requirement'])} | {read} | {why} |")
        lines += ["", "The BoC was renumbered after the deletions (an extractor numbers what it found), so the "
                  "char numbers do not give the deleted rows away.", ""]
    lines += [f"## What the app actually did (verified {today})", ""]
    if doc is None:
        lines += ["Not verified in this build (`--no-check`).", ""]
    else:
        extra = [f for f in doc["flags"] if f not in hits.values()]
        lines += [f"- Rules only, model check off: **{len(doc['flags'])} flag{'s' if len(doc['flags']) != 1 else ''}**; "
                  f"planted items matched exactly: **{'yes' if matched else 'NO'}**.",
                  f"- Rows read from the BoC: {doc['boc']['rows']}. Rows placed on the drawing: "
                  f"{doc['stats'].get('rows_located')} of {doc['boc']['rows']}.",
                  f"- Run time: {secs:.1f} s on a laptop CPU. The text read (OCR) of this drawing ships with the app in "
                  "`samples/cache/`, so no sheet is OCR'd again; on the live site add a few seconds for a cold start. "
                  "A drawing with no cached read takes about 6 to 9 s per sheet for OCR."]
        if full_flags is not None:
            lines.append(f"- The same check on the complete BoC (before the misses were planted): "
                         f"{full_flags} flag{'s' if full_flags != 1 else ''}.")
        if extra:
            lines.append("- Flags not planted: " + "; ".join(f"{f['flag_id']} `{_md_cell(f['drawing_text'])}`" for f in extra))
        ns = notices(doc)
        lines += ["- Notices on the results screen: " + ("none." if not ns else ""), *[f"  - \"{n}\"" for n in ns]]
        if any("ambiguous leftover" in n for n in ns) and not any("R8" in f["rule_ids"] for f in doc["flags"]):
            lines.append("  (That is the app's standard wording when the model check is off. In this run the ambiguous "
                         "items scored under the 0.5 flag threshold, so none of them became a flag.)")
        lines.append("")
    lines += ["## Good to know", "",
              "- With **Use the AI** on (and a key), the model also judges ambiguous leftovers, so results can "
              "differ slightly from this key.",
              "- Only the misses listed above were planted. Every other row matches the drawing.",
              *(["- This is the same pair the app loads with **Use sample drawing**."]
                if pair["format"] == "csv_simulated" else []),
              f"- Source of the drawing: NIST MBE PMI validation and conformance testing test cases (public). "
              f"The original file is `samples/{'eval/' if 'eval' in str(pdf) else ''}{pdf.name}` in this repo.", ""]
    return "\n".join(lines), matched


def build(check: bool) -> list[dict]:
    import datetime as dt

    today = dt.date.today().isoformat()
    OUT.mkdir(exist_ok=True)
    summary = []
    for pair in PAIRS:
        spec = _spec(pair["drawing"])
        folder = OUT / pair["folder"]
        folder.mkdir(exist_ok=True)
        shutil.copyfile(spec["pdf"], folder / "drawing.pdf")
        full = full_rows(pair["drawing"])
        rows, records = plant(full, pair["planted"])
        if pair["format"] == "form3":
            boc_name = "boc.xlsx"
            write_form3(folder / boc_name, rows, spec)
        elif pair["format"] == "csv_simulated":
            boc_name = "boc.csv"  # the bundled demo BoC, byte for byte
            shutil.copyfile(SAMPLES / "ftc06_boc_simulated.csv", folder / boc_name)
            with (folder / boc_name).open(encoding="utf-8") as f:
                assert [r["requirement"] for r in csv.DictReader(f)] == [r["requirement"] for r in rows]
        else:
            boc_name = "boc.csv"
            write_csv(folder / boc_name, rows)
        doc = secs = None
        full_flags = None
        if check:
            doc, secs = check_pair(folder / "drawing.pdf", folder / boc_name)
            if records:  # the same check on the complete BoC, before anything was planted
                import tempfile

                with tempfile.TemporaryDirectory() as td:
                    fp = Path(td) / f"complete{Path(boc_name).suffix}"
                    write_form3(fp, full, spec) if pair["format"] == "form3" else write_csv(fp, full)
                    full_flags = len(check_pair(folder / "drawing.pdf", fp)[0]["flags"])
        text, matched = answer_key(pair, spec, rows, full, records, boc_name, doc, secs, today, full_flags)
        (folder / "ANSWER-KEY.md").write_text(text, encoding="utf-8")
        summary.append({"folder": pair["folder"], "flags": len(doc["flags"]) if doc else None,
                        "expected": len(records), "matched": matched, "seconds": round(secs, 2) if secs else None,
                        "notices": notices(doc) if doc else None})
        print(f"{pair['folder']}: {len(rows)} rows, planted {len(records)}, "
              + (f"flags {len(doc['flags'])}, matched {matched}, {secs:.1f}s" if doc else "not checked"))
        if doc:
            for f in doc["flags"]:
                print(f"    {f['flag_id']} s{f['sheet']} {f['zone']} {f['rule_ids']} {f['drawing_text']!r}")
    zpath = OUT / "tallymark-sample-pairs.zip"
    with zipfile.ZipFile(zpath, "w", zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        for pair in PAIRS:
            for p in sorted((OUT / pair["folder"]).iterdir()):
                z.write(p, f"tallymark-sample-pairs/{pair['folder']}/{p.name}")
        z.write(OUT / "README.md", "tallymark-sample-pairs/README.md") if (OUT / "README.md").exists() else None
    print(f"zip: {zpath.stat().st_size // 1024} KB")
    return summary


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--no-check", action="store_true", help="build the files without running the pipeline")
    a = ap.parse_args()
    summary = build(not a.no_check)
    if not a.no_check and not all(s["matched"] for s in summary):
        raise SystemExit("Some pairs do not match their planted items: fix the BoC, never the app.")


if __name__ == "__main__":
    main()
