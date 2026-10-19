"""Bill of Characteristics input: parse, validate, normalize.

The column schema is ours. Extraction tools export BoC/results to Excel with their own
column names, so common aliases are accepted.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field

from .geometry import Box

ALIASES = {
    "char_no": ["char_no", "char no", "char. no.", "char no.", "characteristic number", "characteristic no", "cn", "balloon", "balloon_no", "balloon no"],
    "sheet": ["sheet", "sheet_no", "page"],
    "zone": ["zone", "reference location", "ref_location", "location"],
    "char_type": ["char_type", "type", "characteristic type"],
    "designator": ["designator", "characteristic designator", "class"],
    "requirement": ["requirement", "specification", "spec", "requirement text", "nominal and tolerance"],
    "nominal": ["nominal"],
    "upper_tol": ["upper_tol", "upper tol", "plus tol", "+tol"],
    "lower_tol": ["lower_tol", "lower tol", "minus tol", "-tol"],
    "units": ["units", "unit"],
    "places": ["places", "qty", "quantity", "instances", "n places"],
    # Form 3 "Results" holds measured values, not a count, so only the explicit name maps here
    "results_count": ["results_count"],
    "bbox_x": ["bbox_x"],
    "bbox_y": ["bbox_y"],
    "bbox_w": ["bbox_w"],
    "bbox_h": ["bbox_h"],
    "source_text": ["source_text", "raw text"],
}
CHAR_TYPES = {"dimension", "gdt", "note", "spec", "surface_finish", "thread", "other"}


class BocError(ValueError):
    """User-facing validation error (copy is shown in the drop zone)."""


@dataclass
class Characteristic:
    char_no: str
    sheet: int
    requirement: str
    zone: str | None = None
    char_type: str = "other"
    designator: str | None = None
    places: int = 1
    results_count: int | None = None
    units: str | None = None
    bbox: Box | None = None
    bbox_source: str | None = None  # "export" | "text_match" | None
    source_text: str | None = None
    located_regions: list[str] = field(default_factory=list)

    def to_dict(self) -> dict:
        return {
            "char_no": self.char_no,
            "sheet": self.sheet,
            "zone": self.zone,
            "char_type": self.char_type,
            "designator": self.designator,
            "requirement": self.requirement,
            "places": self.places,
            "results_count": self.results_count,
            "units": self.units,
            "bbox": self.bbox.to_dict() if self.bbox else None,
            "bbox_source": self.bbox_source,
        }


@dataclass
class ParsedBoc:
    rows: list[Characteristic]
    columns: list[str]
    has_bbox: bool
    warnings: list[str]


HEADER_SCAN_ROWS = 25  # AS9102 Form 3 sheets put a title and part-number block above the column headers


def _read_rows(data: bytes, filename: str) -> list[list[str]]:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:  # noqa: BLE001
            raise BocError("This spreadsheet could not be opened. Save it as XLSX or CSV and try again.") from exc
        return [["" if v is None else str(v) for v in r] for r in wb.worksheets[0].iter_rows(values_only=True)]
    if name.endswith(".csv"):
        return list(csv.reader(io.StringIO(data.decode("utf-8-sig", errors="replace"))))
    raise BocError("This file isn't a CSV or XLSX.")


def _norm_header(h: str) -> str:
    """'5. Char. No.' -> 'char. no.', 'Places (qty)' -> 'places', '8. Requirement:' -> 'requirement'."""
    h = re.sub(r"\s+", " ", (h or "").strip().lower())
    h = re.sub(r"^\d+\.\s*", "", h)  # Form 3 field numbers
    h = re.sub(r"\s*\([^)]*\)$", "", h)  # trailing unit / qualifier in parentheses
    return h.rstrip(":").strip()


def _map_columns(header: list[str]) -> dict[str, str]:
    lower = {_norm_header(h): h for h in header if h and h.strip()}
    mapping: dict[str, str] = {}
    for canon, names in ALIASES.items():
        for n in names:
            if n in lower:
                mapping[canon] = lower[n]
                break
    return mapping


def _read_table(data: bytes, filename: str) -> list[dict[str, str]]:
    """Rows as dicts keyed by the header row. The header is the first row (within the first
    HEADER_SCAN_ROWS) that names both a characteristic number and a requirement."""
    rows = _read_rows(data, filename)
    if not rows:
        return []
    hdr_idx = 0
    for i, r in enumerate(rows[:HEADER_SCAN_ROWS]):
        m = _map_columns(r)
        if "char_no" in m and "requirement" in m:
            hdr_idx = i
            break
    header = [h.strip() for h in rows[hdr_idx]]
    return [
        {header[i]: v for i, v in enumerate(r) if i < len(header) and header[i]}
        for r in rows[hdr_idx + 1:]
        if any((v or "").strip() for v in r)
    ]


_SHEET_IN_LOCATION = re.compile(r"\b(?:sheet|sht|page|pg)\.?\s*(\d+)", re.I)
_ZONE_IN_LOCATION = re.compile(r"\bzone\s*([A-Z]{1,2}\d{1,2})\b", re.I)


def _to_int(v: str | None, default: int | None) -> int | None:
    if v is None or str(v).strip() == "":
        return default
    try:
        return int(float(str(v).strip()))
    except ValueError:
        return default


def _to_float(v: str | None) -> float | None:
    if v is None or str(v).strip() == "":
        return None
    try:
        return float(str(v).strip())
    except ValueError:
        return None


def parse_boc(data: bytes, filename: str) -> ParsedBoc:
    table = _read_table(data, filename)
    if not table:
        raise BocError("This file has no rows.")
    header = list(table[0].keys())
    cols = _map_columns(header)
    # Form 3 has no sheet column: the sheet lives in Reference Location ("Sheet 2, zone B3")
    sheet_from_location = "sheet" not in cols and "zone" in cols and any(
        _SHEET_IN_LOCATION.search(r.get(cols["zone"]) or "") for r in table)
    missing = [c for c in ("char_no", "sheet", "requirement") if c not in cols
               and not (c == "sheet" and sheet_from_location)]
    expected = "Expected a characteristic number, a sheet (or a Reference Location like 'Sheet 1, zone C7') and a requirement."
    if "requirement" in missing:
        raise BocError(f"This file has no requirement column. {expected}")
    if missing:
        raise BocError(f"This file has no {missing[0].replace('_', ' ')} column. {expected}")

    def get(row: dict, canon: str) -> str | None:
        key = cols.get(canon)
        return row.get(key) if key else None

    has_bbox_cols = all(k in cols for k in ("bbox_x", "bbox_y", "bbox_w", "bbox_h"))
    rows: list[Characteristic] = []
    warnings: list[str] = []
    for i, r in enumerate(table, start=2):
        char_no = (get(r, "char_no") or "").strip()
        requirement = (get(r, "requirement") or "").strip()
        location = (get(r, "zone") or "").strip()
        if sheet_from_location:
            m = _SHEET_IN_LOCATION.search(location)
            sheet = int(m.group(1)) if m else None
        else:
            sheet = _to_int(get(r, "sheet"), None)
        zm = _ZONE_IN_LOCATION.search(location)
        zone = zm.group(1).upper() if zm else (location or None)
        if not char_no or not requirement or sheet is None:
            warnings.append(f"Row {i} skipped: missing char_no, sheet or requirement.")
            continue
        bbox = None
        if has_bbox_cols:
            vals = [_to_float(get(r, k)) for k in ("bbox_x", "bbox_y", "bbox_w", "bbox_h")]
            if all(v is not None for v in vals):
                bbox = Box(*vals)  # type: ignore[arg-type]
        ctype = (get(r, "char_type") or "other").strip().lower() or "other"
        rows.append(
            Characteristic(
                char_no=char_no,
                sheet=sheet,
                requirement=requirement,
                zone=zone,
                char_type=ctype if ctype in CHAR_TYPES else "other",
                designator=(get(r, "designator") or None),
                places=_to_int(get(r, "places"), 1) or 1,
                results_count=_to_int(get(r, "results_count"), None),
                units=(get(r, "units") or None),
                bbox=bbox,
                bbox_source="export" if bbox else None,
                source_text=get(r, "source_text"),
            )
        )
    if not rows:
        raise BocError("No usable rows. Each row needs char_no, sheet and requirement.")
    return ParsedBoc(rows=rows, columns=header, has_bbox=any(r.bbox for r in rows), warnings=warnings)


_SYMBOLS = [
    ("±", "+-"), ("+/-", "+-"), ("+ /-", "+-"), ("�", "+-"),
    ("Ø", "0"), ("⌀", "0"), ("Φ", "0"), ("φ", "0"), ("DIA", "0"),
    ("⌖", ""), ("中", ""), ("⌴", ""), ("└┘", ""), ("⏥", ""), ("⌓", ""), ("▽", "V"), ("↧", "V"),
    ("Ⓜ", "M"), ("@", "M"),
]


def normalize_text(text: str) -> str:
    """Fold OCR and typography variants so requirement text and drawing text compare."""
    t = text.upper()
    for a, b in _SYMBOLS:
        t = t.replace(a, b)
    t = re.sub(r"[|/\\\[\]{}]", "", t)
    t = re.sub(r"\s+", "", t)
    t = re.sub(r"(?<![0-9])[0OQ]+\.(\d)", r".\1", t)  # Ø.250 / 0.250 / OCR 'Q.250' -> .250
    t = re.sub(r"L[1J|](?=[.\d])", "", t)  # counterbore symbol read as 'L1' / 'LJ'
    return t
